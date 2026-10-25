"""Read contract text from PDF, DOCX and plain-text uploads."""
from __future__ import annotations

import io
from pathlib import PurePath

SUPPORTED_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx", ".txt")
MAX_BYTES = 20 * 1024 * 1024  # 20 MB safety limit


class DocumentError(Exception):
    """Base class for problems reading a document."""


class UnsupportedDocumentError(DocumentError):
    """Raised for file types ClauseGuard cannot read."""


def normalize_text(text: str) -> str:
    """Normalise line endings and drop NUL characters (offsets are relative to this text)."""
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")


def _read_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("latin-1", errors="replace")


def _read_pdf(data: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("pdfplumber is not installed. Run: pip install pdfplumber") from exc
    pages: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except Exception as exc:
        raise DocumentError(f"Could not read this PDF ({exc.__class__.__name__}). It may be corrupted or encrypted.") from exc
    text = "\n\n".join(p for p in pages if p.strip())
    if not text.strip():
        raise DocumentError("No selectable text found in this PDF. Scanned PDFs need OCR, which ClauseGuard does not include.")
    return text


def _read_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("python-docx is not installed. Run: pip install python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"Could not read this DOCX file ({exc.__class__.__name__}).") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    if not text.strip():
        raise DocumentError("This DOCX file contains no text.")
    return text


def extract_text(data: bytes, filename: str) -> str:
    """Extract normalised text from uploaded bytes based on the file extension."""
    if len(data) > MAX_BYTES:
        raise DocumentError("File is larger than 20 MB.")
    ext = PurePath(filename).suffix.lower()
    if ext == ".txt":
        text = _read_txt(data)
    elif ext == ".pdf":
        text = _read_pdf(data)
    elif ext == ".docx":
        text = _read_docx(data)
    else:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{ext or filename}'. Please upload PDF, DOCX or TXT (old .doc files must be saved as .docx)."
        )
    text = normalize_text(text)
    if not text.strip():
        raise DocumentError("The document is empty.")
    return text
